import io
import re
from typing import Dict, Any
from pptx import Presentation
from schemas import PDFInfoResponse

def extract_text_from_pptx(file_content: bytes) -> Dict[str, Any]:
    """
    從 PPTX 內容中提取標題、內文和頁數。
    """
    try:
        prs = Presentation(io.BytesIO(file_content))
        data = {
            "title": None,
            "abstract": "",
            "page_count": 0,
            "extracted_text_snippet": ""
        }

        # 1. 取得頁數
        data["page_count"] = len(prs.slides)

        full_text = []
        
        # 2. 遍歷每一頁投影片
        for i, slide in enumerate(prs.slides):
            slide_text = []
            
            # 嘗試抓取標題 (通常在第一頁的第一個 shape，或者有 title placeholder)
            if i == 0:
                if slide.shapes.title:
                    data["title"] = slide.shapes.title.text
                else:
                    # 如果沒有明確的 title shape，嘗試找第一個有文字的 shape
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            data["title"] = shape.text.strip()
                            break
            
            # 提取該頁所有文字
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            
            if slide_text:
                full_text.append("\n".join(slide_text))

        # 3. 組合內文作為摘要 (Abstract)
        # 用換行符號連接所有頁面的文字
        all_text_content = "\n\n".join(full_text)
        data["abstract"] = all_text_content
        
        # 4. 擷取片段供除錯/預覽
        data["extracted_text_snippet"] = all_text_content[:500] + "..." if len(all_text_content) > 500 else all_text_content

        return data

    except Exception as e:
        print(f"Error extracting text from PPTX: {e}")
        return {"extracted_text_snippet": f"解析失敗: {str(e)}"}

async def parse_pptx_for_metadata(file_content: bytes) -> PDFInfoResponse:
    """
    解析 PPTX 入口點，回傳統一格式
    """
    metadata = extract_text_from_pptx(file_content)
    return PDFInfoResponse(**metadata)